#!/usr/bin/env python3
"""Erzeugt/aktualisiert die Excel-Uebersicht aller muki-Faelle.

Quelle sind alle ``falldaten.json`` unter ``faelle/`` (ausser ``vorlage``).
Die Datei ``out/Uebersicht_muki.xlsx`` wird bei jedem Lauf neu geschrieben.
Manuelle Spalten (muki-Erstattung, Status, Notizen) werden aus einer bereits
vorhandenen Tabelle uebernommen, damit sie nicht verloren gehen.
"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
FAELLE = BASE / "faelle"
OUT = BASE / "out" / "Uebersicht_muki.xlsx"

HEADERS = [
    "Fall",
    "Datum (Einreichung)",
    "Behandler",
    "Leistung",
    "Honorarnote",
    "Betrag (€)",
    "Diagnose",
    "ÖGK-Erstattung (€)",
    "muki-Erstattung (€)",
    "Status",
    "Notizen",
]
MANUAL = ["muki-Erstattung (€)", "Status", "Notizen"]
# Optionaler Anfangsstatus pro Fall (Schluessel = Fall-Label, z. B.
# "Mustermann/08_26"). Leer lassen; der Status kommt sonst aus der bestehenden
# Tabelle bzw. ist "vorbereitet".
DEFAULT_STATUS: dict[str, str] = {}

CURRENCY = '#,##0.00 "€"'


def money(text: str) -> float | None:
    m = re.search(r"EUR\s*([\d.]+,\d{2})", text or "")
    if not m:
        return None
    return float(m.group(1).replace(".", "").replace(",", "."))


def honorar_nr(text: str) -> str:
    m = re.search(r"Nr\.\s*([^\s,]+)", text or "")
    return m.group(1) if m else ""


def behandler(text: str) -> str:
    m = re.search(r"\(([^)]+)\)", text or "")
    return m.group(1) if m else ""


def fmt_datum(value: str) -> str:
    for f in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(value, f).strftime("%d.%m.%Y")
        except (ValueError, TypeError):
            continue
    return value or ""


def fall_label(path: Path) -> str:
    rel = path.relative_to(FAELLE)
    parts = rel.parts
    if len(parts) == 2:
        data = json.loads(path.read_text(encoding="utf-8"))
        rechnungen = data.get("rechnungen") or []
        sub = Path(rechnungen[0]).parts[0] if rechnungen and "/" in rechnungen[0] else ""
        return f"{parts[0]}/{sub}" if sub else parts[0]
    return "/".join(parts[:-1])


def discover_cases() -> list[Path]:
    return sorted(
        p
        for p in FAELLE.glob("**/falldaten.json")
        if "vorlage" not in p.parts
    )


def load_manual() -> dict[str, dict]:
    if not OUT.exists():
        return {}
    wb = load_workbook(OUT)
    ws = wb.active
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        fall = row[idx["Fall"]]
        if not fall:
            continue
        out[fall] = {h: row[idx[h]] for h in MANUAL if h in idx}
    return out


def build_rows() -> list[dict]:
    manual = load_manual()
    rows = []
    for path in discover_cases():
        data = json.loads(path.read_text(encoding="utf-8"))
        fall = fall_label(path)
        leistungen = data.get("leistungen") or []
        if isinstance(leistungen, str):
            leistungen = [leistungen]
        honorar_text = next((l for l in leistungen if "Honorar" in l), "")
        unterlagen = data.get("unterlagen") or []
        if isinstance(unterlagen, str):
            unterlagen = [unterlagen]
        oegk_text = next(
            (u for u in unterlagen if "ÖGK" in u or "OGK" in u), ""
        )
        m = manual.get(fall, {})
        rows.append(
            {
                "Fall": fall,
                "Datum (Einreichung)": fmt_datum(data.get("datum", "")),
                "Behandler": behandler(honorar_text),
                "Leistung": leistungen[0] if leistungen else "",
                "Honorarnote": honorar_nr(honorar_text),
                "Betrag (€)": money(honorar_text),
                "Diagnose": ", ".join(data.get("diagnose") or []),
                "ÖGK-Erstattung (€)": money(oegk_text),
                "muki-Erstattung (€)": m.get("muki-Erstattung (€)"),
                "Status": m.get("Status") or DEFAULT_STATUS.get(fall, "vorbereitet"),
                "Notizen": m.get("Notizen") or "",
            }
        )
    rows.sort(key=lambda r: datetime.strptime(r["Datum (Einreichung)"], "%d.%m.%Y"))
    return rows


def write_xlsx(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "muki-Fälle"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in rows:
        ws.append([row[h] for h in HEADERS])

    money_cols = {
        HEADERS.index("Betrag (€)") + 1,
        HEADERS.index("ÖGK-Erstattung (€)") + 1,
        HEADERS.index("muki-Erstattung (€)") + 1,
    }
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if c in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = CURRENCY

    widths = [22, 18, 20, 24, 14, 12, 12, 18, 18, 14, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)


def main() -> None:
    rows = build_rows()
    write_xlsx(rows)
    print(f"{len(rows)} Fall/Fälle geschrieben nach:\n  {OUT}")
    for r in rows:
        print(f"  - {r['Fall']}: {r['Behandler']} | {r['Betrag (€)']} € "
              f"| ÖGK {r['ÖGK-Erstattung (€)']} € | {r['Status']}")


if __name__ == "__main__":
    main()
